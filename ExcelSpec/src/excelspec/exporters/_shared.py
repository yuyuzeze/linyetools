"""Shared rendering helpers for DocumentIR exporters."""

from __future__ import annotations

import html
import json
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Literal
from urllib.parse import urlparse

from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries

from ..models.document_ir import (
    AssetIR,
    CellIR,
    DocumentIR,
    RegionIR,
    RegionType,
    SheetIR,
    TableIR,
)


def cell_text(cell: CellIR | None) -> str:
    if cell is None:
        return ""
    if cell.display_value is not None:
        return cell.display_value
    if cell.raw_value is None:
        return ""
    if isinstance(cell.raw_value, (dict, list)):
        return json.dumps(cell.raw_value, ensure_ascii=False, sort_keys=True)
    return str(cell.raw_value)


def table_bounds(table: TableIR) -> tuple[int, int, int, int] | None:
    if not table.cells:
        return None
    return (
        min(cell.row for cell in table.cells),
        max(cell.row + cell.row_span - 1 for cell in table.cells),
        min(cell.column for cell in table.cells),
        max(cell.column + cell.col_span - 1 for cell in table.cells),
    )


def table_columns(table: TableIR) -> list[int]:
    """Return present columns only so trimmed blank padding columns stay removed."""
    return sorted({cell.column for cell in table.cells})


def table_rows(table: TableIR) -> list[int]:
    return sorted({cell.row for cell in table.cells})


def table_cell_map(table: TableIR) -> dict[tuple[int, int], CellIR]:
    return {(cell.row, cell.column): cell for cell in table.cells}


def logical_cell_map(table: TableIR) -> dict[tuple[int, int], CellIR]:
    """Map every represented position to its logical merge master."""
    cells = table_cell_map(table)
    by_coordinate = {cell.coordinate: cell for cell in table.cells}
    result: dict[tuple[int, int], CellIR] = {}
    for position, cell in cells.items():
        if cell.merged_master:
            result[position] = by_coordinate.get(cell.merged_master, cell)
        else:
            result[position] = cell
            for row in range(cell.row, cell.row + cell.row_span):
                for column in range(cell.column, cell.column + cell.col_span):
                    result.setdefault((row, column), cell)
    return result


def has_complex_merges(table: TableIR) -> bool:
    return any(
        cell.merged_master is not None or cell.row_span > 1 or cell.col_span > 1
        for cell in table.cells
    )


def render_html_table(table: TableIR, *, css_class: str = "") -> str:
    bounds = table_bounds(table)
    class_attribute = f' class="{html.escape(css_class, quote=True)}"' if css_class else ""
    if bounds is None:
        return f"<table{class_attribute}></table>"

    min_row, max_row, _, _ = bounds
    columns = table_columns(table)
    cells = table_cell_map(table)
    covered = {
        (row, column)
        for cell in table.cells
        if not cell.merged_master
        for row in range(cell.row, cell.row + cell.row_span)
        for column in range(cell.column, cell.column + cell.col_span)
        if (row, column) != (cell.row, cell.column)
    }
    lines = [f"<table{class_attribute}>"]
    for row in range(min_row, max_row + 1):
        if not any((row, column) in cells or (row, column) in covered for column in columns):
            continue
        lines.append("  <tr>")
        for column in columns:
            cell = cells.get((row, column))
            if (row, column) in covered or (cell is not None and cell.merged_master):
                continue
            tag = "th" if row < min_row + table.header_rows else "td"
            attributes: list[str] = []
            if cell is not None and cell.row_span > 1:
                attributes.append(f'rowspan="{cell.row_span}"')
            if cell is not None and cell.col_span > 1:
                attributes.append(f'colspan="{cell.col_span}"')
            if cell is not None:
                attributes.append(f'data-source-cell="{html.escape(cell.coordinate)}"')
            attribute_text = f" {' '.join(attributes)}" if attributes else ""
            value = html.escape(cell_text(cell)).replace("\n", "<br>")
            lines.append(f"    <{tag}{attribute_text}>{value}</{tag}>")
        lines.append("  </tr>")
    lines.append("</table>")
    return "\n".join(lines)


def all_assets(document: DocumentIR, sheet: SheetIR | None = None) -> dict[str, AssetIR]:
    assets = {asset.asset_id: asset for asset in document.assets}
    if sheet is not None:
        assets.update({asset.asset_id: asset for asset in sheet.assets})
    return assets


def asset_uri(asset: AssetIR, destination: Path) -> str:
    if "://" in asset.uri and not asset.uri.startswith("file://"):
        return asset.uri
    parsed = urlparse(asset.uri)
    raw_path = Path(parsed.path if parsed.scheme == "file" else asset.uri)
    if not raw_path.is_absolute():
        return raw_path.as_posix()
    try:
        return Path(os.path.relpath(raw_path, destination.parent)).as_posix()
    except ValueError:
        return raw_path.as_uri()


# --- Readable (MD/HTML) rendering helpers -----------------------------------
#
# Templates for 方眼 (graph-paper) style Excel sheets often spread one
# business field across several merged physical columns, and headers can be
# split across the multiple `header_rows` of a merged group. The helpers
# below give exporters a business-language view: only the columns a template
# actually mapped via `column_semantics`, with merge text already resolved,
# so a compact table (or a plain list for header_rows == 0 captures) can be
# rendered without hard-coding any sheet's physical coordinates. They are
# purely additive -- existing helpers above (table_columns/table_cell_map/
# cell_text/logical_cell_map) are untouched so the canonical JSON/JSONL
# exporters keep their current output.


def semantic_columns(table: TableIR) -> list[int]:
    """Physical columns to show for a readable rendering.

    Tables whose template mapped `column_semantics` only surface those
    columns, left-to-right; a 方眼 grid's extra physical/merge-member
    columns are dropped. Tables without a mapping (freeform captures) fall
    back to every present column so nothing is silently lost.
    """
    present = table_columns(table)
    if not table.column_semantics:
        return present
    present_set = set(present)
    indices = sorted(
        {
            column_index_from_string(letter)
            for letter in table.column_semantics
            if column_index_from_string(letter) in present_set
        }
    )
    return indices or present


def header_label_row(table: TableIR) -> list[str]:
    """Business-language header text for `semantic_columns(table)`."""
    columns = semantic_columns(table)
    labels = table.metadata.get("header_labels")
    if not isinstance(labels, dict):
        return ["" for _ in columns]
    return [str(labels.get(get_column_letter(column), "")) for column in columns]


def compact_rows(table: TableIR) -> list[list[str]]:
    """Row-major cell text for `semantic_columns(table)`, header rows
    excluded and merge text resolved via `logical_cell_map`. Blank rows are
    dropped. Also doubles as the "does this table have real content" check.
    """
    bounds = table_bounds(table)
    if bounds is None:
        return []
    min_row, max_row, _, _ = bounds
    columns = semantic_columns(table)
    cells = logical_cell_map(table)
    start = min_row + max(table.header_rows, 0)
    rows: list[list[str]] = []
    for row in range(start, max_row + 1):
        values = [cell_text(cells.get((row, column))) for column in columns]
        if not any(values):
            continue
        rows.append(_normalize_group_row(values))
    return rows


def is_group_row(values: list[str]) -> bool:
    """Detect section banners like 「グループ：基本情報入力」."""
    non_empty = [value.strip() for value in values if value and value.strip()]
    if not non_empty:
        return False
    text = non_empty[0]
    if not re.match(r"^(グループ|Group)\s*[:：]", text, flags=re.IGNORECASE):
        return False
    return len(set(non_empty)) == 1


def _normalize_group_row(values: list[str]) -> list[str]:
    if not is_group_row(values):
        return values
    text = next(value.strip() for value in values if value and value.strip())
    return [text if index == 0 else "" for index, _ in enumerate(values)]


def compact_list_rows(table: TableIR) -> list[list[str]]:
    """Like `compact_rows`, but for the header_rows == 0 plain-list
    rendering: consecutive semantic columns resolving to the very same
    merged cell (e.g. a banner-style label spread across many physical
    columns) collapse into a single value instead of repeating. Table
    rendering keeps `compact_rows` -- it needs one aligned cell per
    semantic column to stay a valid grid.
    """
    return [values for _, values in compact_list_rows_with_positions(table)]


def compact_list_rows_with_positions(table: TableIR) -> list[tuple[int, list[str]]]:
    """Return ``(excel_row, values)`` pairs for non-empty freeform rows."""
    bounds = table_bounds(table)
    if bounds is None:
        return []
    min_row, max_row, _, _ = bounds
    columns = semantic_columns(table)
    cells = logical_cell_map(table)
    start = min_row + max(table.header_rows, 0)
    rows: list[tuple[int, list[str]]] = []
    for row in range(start, max_row + 1):
        values: list[str] = []
        last_identity: str | None = None
        for column in columns:
            cell = cells.get((row, column))
            identity = cell.coordinate if cell is not None else None
            if identity is not None and identity == last_identity:
                last_identity = identity
                continue
            last_identity = identity
            text = cell_text(cell)
            if text:
                values.append(text)
        if values:
            rows.append((row, values))
    return rows


def asset_anchor_row(asset: AssetIR) -> int | None:
    """Best-effort top Excel row for an asset anchor / source range."""
    reference = asset.anchor
    if reference is None and asset.source is not None:
        reference = asset.source.range or asset.source.cell
    if not reference:
        return None
    try:
        if ":" in reference:
            return range_boundaries(reference)[1]
        return range_boundaries(f"{reference}:{reference}")[1]
    except ValueError:
        match = re.search(r"(\d+)", reference)
        return int(match.group(1)) if match else None


@dataclass(frozen=True, slots=True)
class RegionBlock:
    """One readable fragment inside a region, ordered by sheet row."""

    kind: Literal["text", "asset"]
    row: int
    text: str | None = None
    asset_id: str | None = None


def should_interleave_region(region: RegionIR) -> bool:
    """Layout/mockup regions keep captions and images in vertical order."""
    if region.metadata.get("readable_mode") == "screenshot":
        return False
    return (
        region.region_type == RegionType.LAYOUT
        or region.metadata.get("extractor_kind") == "asset"
    )


def interleaved_region_blocks(
    region: RegionIR, assets: dict[str, AssetIR]
) -> list[RegionBlock]:
    """Merge freeform text rows and region assets by Excel row order.

    Text on the same row as an image is emitted first so captions above an
    image stay above it, and notes below follow after the image row.
    """
    events: list[tuple[int, int, int, RegionBlock]] = []
    sequence = 0
    for table in region.tables:
        if table.header_rows > 0:
            continue
        for row, values in compact_list_rows_with_positions(table):
            text = " / ".join(values)
            events.append(
                (
                    row,
                    0,
                    sequence,
                    RegionBlock(kind="text", row=row, text=text),
                )
            )
            sequence += 1
    for index, asset_id in enumerate(region.asset_ids):
        if asset_id not in assets:
            continue
        row = asset_anchor_row(assets[asset_id])
        # Unanchored assets keep relative declaration order after known rows.
        sort_row = row if row is not None else 10**9
        events.append(
            (
                sort_row,
                1,
                index,
                RegionBlock(kind="asset", row=sort_row, asset_id=asset_id),
            )
        )
    events.sort(key=lambda item: (item[0], item[1], item[2]))
    return [item[3] for item in events]


def table_has_content(table: TableIR) -> bool:
    return bool(compact_rows(table))


def is_unrecognized_region(region: RegionIR) -> bool:
    return region.region_id.startswith("unrecognized")


def _is_empty_value(value: object) -> bool:
    return value is None or value == "" or value == [] or value == {}


def region_has_readable_content(region: RegionIR) -> bool:
    if any(not _is_empty_value(value) for value in region.values.values()):
        return True
    if region.asset_ids:
        return True
    return any(table_has_content(table) for table in region.tables)


def should_render_region(region: RegionIR) -> bool:
    """MD/HTML should skip parser noise: unrecognized-* freeform captures
    and regions with no non-empty values, no valid table content, and no
    assets. Canonical JSON/JSONL keep every region untouched."""
    return not is_unrecognized_region(region) and region_has_readable_content(region)


def readable_document_metadata(document: DocumentIR) -> list[tuple[str, object]]:
    """MD/HTML are business-readable views, so omit all parser metadata.

    The title and extracted region values remain visible. Canonical JSON/JSONL
    retain document identity, template selection, source and diagnostics.
    """
    return []


def readable_region_values(region: RegionIR) -> list[tuple[str, object]]:
    """Return populated values for MD/HTML using Excel labels when known.

    IR keeps machine keys from ``key_semantics`` (e.g. ``document_no``).
    Readable exporters reverse ``metadata.key_labels`` so MD shows
    「文書番号」 instead of ``document_no``.
    """
    labels = region.metadata.get("key_labels")
    reverse: dict[str, str] = {}
    if isinstance(labels, dict):
        for label, semantic in labels.items():
            if isinstance(label, str) and isinstance(semantic, str) and semantic:
                reverse.setdefault(semantic, label)
    return [
        (reverse.get(key, key), value)
        for key, value in region.values.items()
        if not _is_empty_value(value)
    ]
