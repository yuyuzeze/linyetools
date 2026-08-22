"""Route CandidateRegions into canonical DocumentIR RegionIR objects.

Each candidate is materialised **only over its own bounds** (via
``ingest.adapter.materialize_region``) before routing — never the whole sheet.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl.utils import get_column_letter

from ..ingest.adapter import materialize_region
from ..ingest.sparse_model import SparseSheet, SparseWorkbookIR
from ..models.document_ir import (
    CellIR,
    DiagnosticIR,
    DiagnosticSeverity,
    RegionIR,
    SourceRef,
    StyleIR,
    TableIR,
)
from .models import CandidateRegion, CandidateRegionType


def _display(cell: CellIR | None) -> str:
    if cell is None:
        return ""
    if cell.display_value is not None:
        return str(cell.display_value).strip()
    if cell.raw_value is not None:
        return str(cell.raw_value).strip()
    return ""


def _is_headery(cell: CellIR | None, styles: dict[int, StyleIR]) -> bool:
    if cell is None or not _display(cell):
        return False
    # style is already resolved onto the CellIR by materialize_region
    style = cell.style
    return bool(style and (style.font.get("bold") or _has_fill(style)))


def _has_fill(style: StyleIR | None) -> bool:
    if not style or not style.fill:
        return False
    fill_type = style.fill.get("type")
    return bool(fill_type) and fill_type != "none"


def _is_numeric(cell: CellIR | None) -> bool:
    if cell is None:
        return False
    if cell.data_type == "n" and isinstance(cell.raw_value, (int, float)) and not isinstance(cell.raw_value, bool):
        return True
    text = _display(cell)
    if not text:
        return False
    try:
        float(text.replace(",", ""))
        return True
    except ValueError:
        return False


def _row_signals(cells_by_pos, row, min_col, max_col, styles) -> dict:
    row_cells = [cells_by_pos.get((row, col)) for col in range(min_col, max_col + 1)]
    present = [c for c in row_cells if c is not None and _display(c)]
    merged = [
        c for c in row_cells
        if c is not None and (c.row_span > 1 or c.col_span > 1 or c.merged_master)
    ]
    styled = sum(_is_headery(c, styles) for c in present)
    numeric = sum(_is_numeric(c) for c in present)
    return {
        "present": len(present),
        "styled_frac": styled / len(present) if present else 0.0,
        "numeric_frac": numeric / len(present) if present else 0.0,
        "has_merge": bool(merged),
    }


def detect_header_rows(cells: list[CellIR], bounds, styles) -> tuple[int, dict]:
    """Decide header row count from styling + merges + data-type transition.

    Returns ``(header_rows, decision)`` where ``decision`` carries the evidence
    and a confidence, so a low-confidence choice is observable and conservative.
    """

    min_row, min_col, max_row, max_col = bounds
    by_position = {(c.row, c.column): c for c in cells}
    evidence: list[str] = []
    header_rows = 0
    window_end = min(max_row, min_row + 3)
    for row in range(min_row, window_end + 1):
        signals = _row_signals(by_position, row, min_col, max_col, styles)
        if signals["present"] == 0:
            break
        header_like = (signals["styled_frac"] >= 0.5 or signals["has_merge"]) and signals[
            "numeric_frac"
        ] < 0.5
        data_like = signals["numeric_frac"] >= 0.5 and signals["styled_frac"] < 0.5
        if header_like and not (data_like and row > min_row):
            header_rows += 1
            if signals["has_merge"] and "merged-header" not in evidence:
                evidence.append("merged-header")
            if signals["styled_frac"] >= 0.5 and "styled-row" not in evidence:
                evidence.append("styled-row")
        else:
            if data_like:
                evidence.append(f"data-type-transition@row{row}")
            break

    header_rows = max(1, header_rows)
    # confidence: a clear styled/merged header with a data row after it is strong;
    # a bare first-row-only guess is weak.
    if len(evidence) >= 2 or (header_rows >= 2 and evidence):
        confidence = 0.9
    elif evidence:
        confidence = 0.7
    else:
        confidence = 0.4
        evidence.append("default-first-row")
    return header_rows, {
        "header_rows": header_rows,
        "confidence": round(confidence, 4),
        "evidence": evidence,
    }


def _header_labels(cells: list[CellIR], bounds, header_rows: int) -> dict[int, str]:
    min_row, min_col, max_row, max_col = bounds
    by_coordinate = {c.coordinate: c for c in cells}
    values: dict[tuple[int, int], str] = {}
    for cell in cells:
        value = _display(cell)
        if not value and cell.merged_master:
            value = _display(by_coordinate.get(cell.merged_master))
        values[(cell.row, cell.column)] = value
    labels: dict[int, str] = {}
    for column in range(min_col, max_col + 1):
        parts = [
            values.get((row, column), "")
            for row in range(min_row, min_row + header_rows)
            if values.get((row, column), "")
        ]
        labels[column] = " / ".join(dict.fromkeys(parts))
    return labels


class RegionRouter:
    def __init__(self, workbook: SparseWorkbookIR) -> None:
        self.workbook = workbook
        self.path = Path(workbook.path)
        self.styles = workbook.styles

    def _materialize(self, sheet: SparseSheet, candidate: CandidateRegion):
        return materialize_region(
            sheet, candidate.bounds.as_tuple(), self.styles, self.path
        )

    def route(
        self, sheet: SparseSheet, candidate: CandidateRegion
    ) -> RegionIR:
        source = SourceRef(
            sheet=sheet.name,
            range=candidate.bounds.a1(),
            workbook_path=str(self.path),
        )
        region = RegionIR(
            region_id=candidate.region_id,
            region_type=candidate.region_type.to_region_type(),
            title=candidate.title,
            source=source,
            confidence=candidate.confidence,
            asset_ids=list(candidate.asset_refs),
            metadata={
                "detection_method": candidate.detection_method,
                "candidate_type": candidate.region_type.value,
                "features": candidate.features,
            },
        )
        if candidate.title_cell:
            region.metadata["title_range"] = candidate.title_cell
        for diagnostic in candidate.diagnostics:
            region.metadata.setdefault("diagnostics", []).append(diagnostic.to_dict())

        handler = {
            CandidateRegionType.TABLE: self._route_table,
            CandidateRegionType.KEY_VALUE: self._route_key_value,
            CandidateRegionType.TEXT: self._route_text,
            CandidateRegionType.IMAGE: self._route_asset,
            CandidateRegionType.SHAPE: self._route_asset,
            CandidateRegionType.LAYOUT: self._route_layout,
            CandidateRegionType.FREEFORM: self._route_text,
        }[candidate.region_type]
        handler(sheet, candidate, region)
        return region

    # -- per-type routers ------------------------------------------------------

    def _route_table(self, sheet, candidate, region) -> None:
        cells, diagnostics = self._materialize(sheet, candidate)
        region.metadata["materialized_cell_count"] = len(cells)
        bounds = candidate.bounds.as_tuple()
        header_rows, decision = detect_header_rows(cells, bounds, self.styles)
        region.metadata["header_decision"] = decision
        if decision["confidence"] < 0.55:
            region.metadata.setdefault("diagnostics", []).append(
                DiagnosticIR(
                    code="route.low_confidence_header",
                    severity=DiagnosticSeverity.INFO,
                    message=(
                        f"表头行数低置信度，保守取 {header_rows} 行: {candidate.region_id}"
                    ),
                    source=region.source,
                    region_id=candidate.region_id,
                ).to_dict()
            )
        labels = _header_labels(cells, bounds, header_rows)
        region.tables.append(
            TableIR(
                table_id=candidate.region_id,
                cells=cells,
                source=region.source,
                header_rows=header_rows,
                column_semantics={},
                metadata={
                    "header_labels": {
                        get_column_letter(col): label for col, label in labels.items()
                    }
                },
            )
        )

    def _route_key_value(self, sheet, candidate, region) -> None:
        cells, _ = self._materialize(sheet, candidate)
        region.metadata["materialized_cell_count"] = len(cells)
        by_position = {(c.row, c.column): c for c in cells}
        min_row, min_col, max_row, max_col = candidate.bounds.as_tuple()
        columns = sorted({c.column for c in cells if _display(c)})
        # pair consecutive populated columns: (label, value)[, (label, value)...]
        pairs = [(columns[i], columns[i + 1]) for i in range(0, len(columns) - 1, 2)]
        for row in range(min_row, max_row + 1):
            for label_col, value_col in pairs:
                label = _display(by_position.get((row, label_col)))
                if not label:
                    continue
                value_cell = by_position.get((row, value_col))
                if value_cell is None:
                    region.values.setdefault(label, None)
                elif value_cell.formula is not None:
                    region.values[label] = value_cell.display_value
                else:
                    region.values[label] = (
                        value_cell.raw_value
                        if value_cell.raw_value is not None
                        else _display(value_cell)
                    )
        # keep the raw cells too so nothing is lost
        region.tables.append(
            TableIR(table_id=candidate.region_id, cells=cells, source=region.source)
        )

    def _route_text(self, sheet, candidate, region) -> None:
        cells, _ = self._materialize(sheet, candidate)
        region.metadata["materialized_cell_count"] = len(cells)
        region.tables.append(
            TableIR(table_id=candidate.region_id, cells=cells, source=region.source)
        )

    def _route_asset(self, sheet, candidate, region) -> None:
        # image / shape: no grid cells, the asset carries the content
        region.metadata["materialized_cell_count"] = 0

    def _route_layout(self, sheet, candidate, region) -> None:
        # fast mode: keep structure + drawings + source range only, no COM.
        cells, _ = self._materialize(sheet, candidate)
        region.metadata["materialized_cell_count"] = len(cells)
        region.metadata["visual"] = True
        region.tables.append(
            TableIR(table_id=candidate.region_id, cells=cells, source=region.source)
        )
        region.metadata.setdefault("diagnostics", []).append(
            DiagnosticIR(
                code="route.layout_visual",
                severity=DiagnosticSeverity.INFO,
                message=f"视觉/布局区域 {candidate.region_id}：fast 模式仅保留结构与资源引用",
                source=region.source,
                region_id=candidate.region_id,
            ).to_dict()
        )


def route_candidates(
    workbook: SparseWorkbookIR,
    sheet: SparseSheet,
    candidates: list[CandidateRegion],
) -> list[RegionIR]:
    router = RegionRouter(workbook)
    return [router.route(sheet, candidate) for candidate in candidates]


__all__ = ["RegionRouter", "route_candidates"]
